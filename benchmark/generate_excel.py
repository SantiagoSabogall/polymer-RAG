import json
import os
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

RESULTS_DIR = Path(__file__).resolve().parent.parent / "temp" / "results"
OUTPUT_FILE = Path(__file__).resolve().parent.parent / "temp" / "benchmark_results.xlsx"

MODELS = [
    ("deepseek-v4-flash-0731", "DeepSeek V4 Flash"),
    ("gpt-5.6-luna", "GPT 5.6 Luna"),
    ("gemini-3.8-flash", "Gemini 3.8 Flash"),
    ("glm-5.3-flash", "GLM 5.3 Flash"),
]

HEADERS = ["Polymer", "WVTR Value", "Units", "Temperature", "RH", "Thickness", "Test Method"]

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name="Calibri", bold=True, size=13, color="1F4E79")
META_FONT = Font(name="Calibri", italic=True, size=10, color="666666")
DATA_FONT = Font(name="Calibri", size=10)
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def load_model_results(model_id: str) -> dict:
    results = {}
    for json_file in sorted(RESULTS_DIR.glob(f"*__{model_id}.json")):
        article_name = json_file.stem.replace(f"__{model_id}", "")
        data = json.loads(json_file.read_text())
        results[article_name] = data
    return results


def write_table(ws, start_row, article_name, data):
    registros = data.get("registros", [])
    title = data.get("article_title", article_name)
    meta = data.get("_meta", {})

    ws.cell(row=start_row, column=1, value=f"Artículo: {title}").font = TITLE_FONT
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=7)

    meta_text = f"Latencia: {meta.get('latency', 'N/A')}s | Tokens: {meta.get('tokens_in', 0)} in / {meta.get('tokens_out', 0)} out"
    ws.cell(row=start_row + 1, column=1, value=meta_text).font = META_FONT
    ws.merge_cells(start_row=start_row + 1, start_column=1, end_row=start_row + 1, end_column=7)

    header_row = start_row + 3
    for col, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN_BORDER

    if not registros:
        ws.cell(row=header_row + 1, column=1, value="No WVTR records found").font = META_FONT
        ws.merge_cells(start_row=header_row + 1, start_column=1, end_row=header_row + 1, end_column=7)
        return header_row + 3

    for i, reg in enumerate(registros):
        row = header_row + 1 + i
        values = [
            reg.get("polymer"),
            reg.get("wvtr_value"),
            reg.get("wvtr_units"),
            reg.get("temperature"),
            reg.get("rh"),
            reg.get("thickness"),
            reg.get("test_method"),
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=val if val is not None else "—")
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            cell.alignment = CENTER if col > 1 else LEFT

    return header_row + len(registros) + 2


def create_excel():
    wb = Workbook()
    wb.remove(wb.active)

    for model_id, model_name in MODELS:
        ws = wb.create_sheet(title=model_name[:31])
        ws.sheet_properties.tabColor = "1F4E79"

        for col in range(1, 8):
            ws.column_dimensions[get_column_letter(col)].width = 22

        ws.column_dimensions["A"].width = 35

        results = load_model_results(model_id)

        ws.cell(row=1, column=1, value=f"Benchmark Results: {model_name}").font = Font(
            name="Calibri", bold=True, size=16, color="1F4E79"
        )
        ws.merge_cells("A1:G1")

        ws.cell(row=2, column=1, value=f"Model ID: {model_id}").font = META_FONT
        ws.merge_cells("A2:G2")

        total_records = sum(len(d.get("registros", [])) for d in results.values())
        ws.cell(row=3, column=1, value=f"Total articles: {len(results)} | Total WVTR records: {total_records}").font = META_FONT
        ws.merge_cells("A3:G3")

        current_row = 5
        for article_name in sorted(results.keys()):
            data = results[article_name]
            current_row = write_table(ws, current_row, article_name, data)
            current_row += 2

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(OUTPUT_FILE))
    print(f"Excel guardado en: {OUTPUT_FILE}")
    print(f"4 sheets (modelos) × 7 tablas (artículos)")


if __name__ == "__main__":
    create_excel()
